#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse, hashlib, json, re
from pathlib import Path

import chromadb
import fitz
import pandas as pd
import yaml
from docx import Document
from sentence_transformers import SentenceTransformer

SUPPORTED_EXTS={'.pdf','.docx','.txt','.md','.csv','.xlsx','.yaml','.yml'}

def clean(s):
    s='' if s is None else str(s)
    s=s.replace('\xa0',' ').replace('\u3000',' ').replace('\r','\n')
    s=re.sub(r'[ \t]+',' ',s)
    s=re.sub(r'\n{3,}','\n\n',s)
    return s.strip()

def sid(*parts):
    return hashlib.sha1('||'.join(map(str,parts)).encode('utf-8')).hexdigest()

def chunks(text,size=700,overlap=120):
    text=clean(text)
    if not text:return []
    paras=[x.strip() for x in re.split(r'\n+',text) if x.strip()]
    out=[]; cur=''
    for p in paras:
        if len(cur)+len(p)+1<=size:
            cur=(cur+'\n'+p).strip()
        else:
            if cur: out.append(cur)
            cur=p
    if cur: out.append(cur)
    final=[]
    for x in out:
        if len(x)<=size: final.append(x); continue
        st=0
        while st<len(x):
            ed=min(len(x),st+size)
            final.append(x[st:ed].strip())
            if ed>=len(x): break
            st=max(st+1,ed-overlap)
    return [x for x in final if x]

def extract(path):
    ext=path.suffix.lower(); items=[]
    if ext=='.pdf':
        doc=fitz.open(str(path))
        for i in range(len(doc)):
            t=clean(doc[i].get_text('text'))
            if t: items.append((t,i+1,f'page_{i+1}'))
    elif ext=='.docx':
        doc=Document(str(path))
        body='\n'.join(clean(p.text) for p in doc.paragraphs if clean(p.text))
        if body: items.append((body,0,'paragraphs'))
        for ti,tb in enumerate(doc.tables,1):
            lines=[]
            for row in tb.rows:
                vals=[clean(c.text) for c in row.cells]
                if any(vals): lines.append(' | '.join(vals))
            if lines: items.append(('\n'.join(lines),0,f'table_{ti}'))
    elif ext in {'.txt','.md'}:
        t=clean(path.read_text(encoding='utf-8',errors='ignore'))
        if t: items.append((t,0,'text'))
    elif ext=='.csv':
        df=pd.read_csv(path,dtype=str,keep_default_na=False)
        lines=[' | '.join(f'{c}: {clean(r[c])}' for c in df.columns if clean(r[c])) for _,r in df.iterrows()]
        t='\n'.join(x for x in lines if x)
        if t: items.append((t,0,'csv'))
    elif ext=='.xlsx':
        xls=pd.ExcelFile(path)
        for sh in xls.sheet_names:
            df=pd.read_excel(path,sheet_name=sh,dtype=str,keep_default_na=False)
            lines=[' | '.join(f'{c}: {clean(r[c])}' for c in df.columns if clean(r[c])) for _,r in df.iterrows()]
            t='\n'.join(x for x in lines if x)
            if t: items.append((t,0,f'sheet_{sh}'))
    elif ext in {'.yaml','.yml'}:
        obj=yaml.safe_load(path.read_text(encoding='utf-8',errors='ignore'))
        items.append((yaml.safe_dump(obj,allow_unicode=True,sort_keys=False,width=120),0,'yaml'))
    return items

def discover(project):
    ws=project/'workspace'; rules=project/'rules'; out=[]
    g=ws/'_global'
    if g.exists():
        for p in g.rglob('*'):
            if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS:
                out.append((p,'_GLOBAL_','global'))
    if ws.exists():
        for d in sorted(ws.iterdir()):
            if not d.is_dir() or d.name.startswith('_'): continue
            m=d/'manuals'
            if m.exists():
                for p in m.rglob('*'):
                    if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS:
                        out.append((p,d.name,'manual'))
    if rules.exists():
        for p in rules.rglob('*'):
            if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS:
                out.append((p,'_GLOBAL_','rule'))
    return out

def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--project',type=Path,default=Path.home()/'yolo_device_monitor')
    ap.add_argument('--embedding-model',required=True)
    ap.add_argument('--collection',default='device_manuals')
    ap.add_argument('--chunk-size',type=int,default=700)
    ap.add_argument('--chunk-overlap',type=int,default=120)
    ap.add_argument('--reset',action='store_true')
    a=ap.parse_args()

    project=a.project.expanduser().resolve(); model_path=Path(a.embedding_model).expanduser().resolve()
    rag=project/'rag'; dbdir=rag/'chroma_db'; rag.mkdir(parents=True,exist_ok=True); dbdir.mkdir(parents=True,exist_ok=True)
    if not model_path.exists(): raise SystemExit(f'Embedding模型不存在: {model_path}')

    sources=discover(project)
    if not sources: raise SystemExit('没有发现可用于RAG的文档')
    print(f'发现文档: {len(sources)}')

    model=SentenceTransformer(str(model_path),device='cuda')
    client=chromadb.PersistentClient(path=str(dbdir))
    if a.reset:
        try: client.delete_collection(a.collection)
        except Exception: pass
    col=client.get_or_create_collection(name=a.collection,metadata={'hnsw:space':'cosine'})

    ids=[]; docs=[]; metas=[]; manifest=[]
    for path,device,stype in sources:
        print(f'解析 [{stype}] {device}: {path.name}')
        try: sections=extract(path)
        except Exception as e:
            print('  ERROR:',repr(e)); continue
        for text,page,section in sections:
            for ci,ch in enumerate(chunks(text,a.chunk_size,a.chunk_overlap)):
                did=sid(device,stype,path,section,page,ci,ch)
                meta={'device_name':device,'source_type':stype,'source_file':path.name,'source_path':str(path),'section':section,'page':int(page),'chunk_index':int(ci)}
                ids.append(did); docs.append(ch); metas.append(meta)
                manifest.append({'id':did,**meta,'chars':len(ch),'preview':ch[:120].replace('\n',' ')})

    if not docs: raise SystemExit('没有生成任何chunk')
    print(f'总chunks: {len(docs)}，开始计算向量...')
    emb=model.encode(docs,batch_size=64,show_progress_bar=True,normalize_embeddings=True)
    for st in range(0,len(docs),1000):
        ed=min(len(docs),st+1000)
        col.upsert(ids=ids[st:ed],documents=docs[st:ed],metadatas=metas[st:ed],embeddings=emb[st:ed].tolist())

    pd.DataFrame(manifest).to_csv(rag/'rag_manifest.csv',index=False,encoding='utf-8-sig')
    stats={'collection':a.collection,'source_files':len(sources),'chunks':len(docs),'embedding_model':str(model_path),'db':str(dbdir)}
    (rag/'rag_stats.json').write_text(json.dumps(stats,ensure_ascii=False,indent=2),encoding='utf-8')
    print('\nRAG知识库构建完成')
    print('ChromaDB:',dbdir)
    print('Manifest:',rag/'rag_manifest.csv')
    print('Stats:',rag/'rag_stats.json')

if __name__=='__main__': main()

