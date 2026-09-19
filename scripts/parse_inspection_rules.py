#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
parse_inspection_rules.py

用途：
从“机房设备指示灯巡检总表.docx”中提取设备指示灯巡检规则，
生成：

1. rules/device_indicator_rules.yaml
2. rules/device_indicator_rules.csv
3. rules/device_indicator_rules_review.xlsx
4. rules/parse_summary.txt

适配常见巡检表结构，例如列：
设备名称 | 指示灯 | 颜色 | 状态 | 含义 | 正常状态

特点：
- 支持 Word 表格中的合并单元格
- 自动向下补全设备名称、指示灯名称、颜色等缺失值
- 自动标准化“常亮 / 闪烁 / 熄灭 / 亮 / 灭”等状态
- 自动推断 normal / info / warning / critical / unknown
- 保留原始文本，便于人工校验
- 不修改原始 Word 文件

使用示例：

cd ~/yolo_device_monitor
source .venv/bin/activate

python scripts/parse_inspection_rules.py \
  --input "/home/crscs/桌面/数字孪生机房巡检识别模型测试视频/机房设备指示灯巡检总表.docx"

如缺少依赖：
pip install python-docx pyyaml openpyxl pandas
"""

import argparse
import re
from collections import OrderedDict
from pathlib import Path

import pandas as pd
import yaml
from docx import Document


COLUMN_ALIASES = {
    "device_name": [
        "设备名称", "设备名", "设备", "设备型号", "名称"
    ],
    "indicator_name": [
        "指示灯", "指示灯名称", "灯名称", "灯名", "状态灯", "告警灯"
    ],
    "color": [
        "颜色", "灯色", "指示灯颜色", "颜 色"
    ],
    "state": [
        "状态", "灯状态", "指示灯状态", "状 态"
    ],
    "meaning": [
        "含义", "说明", "状态含义", "指示灯含义", "意义"
    ],
    "normal_state": [
        "正常状态", "正常", "正常情况", "正常指示", "正常状态说明"
    ],
}

STATE_MAP = {
    "常亮": "steady",
    "长亮": "steady",
    "亮": "on",
    "点亮": "on",
    "亮起": "on",
    "闪烁": "blinking",
    "闪": "blinking",
    "慢闪": "slow_blinking",
    "快闪": "fast_blinking",
    "熄灭": "off",
    "灭": "off",
    "不亮": "off",
}

COLOR_MAP = {
    "绿": "green",
    "绿色": "green",
    "红": "red",
    "红色": "red",
    "黄": "yellow",
    "黄色": "yellow",
    "橙": "orange",
    "橙色": "orange",
    "蓝": "blue",
    "蓝色": "blue",
    "白": "white",
    "白色": "white",
    "紫": "purple",
    "紫色": "purple",
    "黄/绿": "yellow_green",
    "绿/黄": "green_yellow",
}

CRITICAL_WORDS = [
    "严重故障", "严重告警", "故障", "异常", "失电", "断电",
    "链路down", "链路 down", "硬件故障", "电源故障", "告警"
]

WARNING_WORDS = [
    "次要告警", "警告", "预警", "注意", "降级", "异常趋势"
]

INFO_WORDS = [
    "启动", "加载", "初始化", "自检", "备用", "待机"
]

NORMAL_WORDS = [
    "正常", "无告警", "供电正常", "链路up", "链路 up",
    "硬件正常", "运行正常", "业务激活", "对接正常"
]


def parse_args():
    p = argparse.ArgumentParser(description="解析 Word 设备指示灯巡检总表")
    p.add_argument(
        "--input",
        type=Path,
        required=True,
        help="巡检总表 .docx 文件路径",
    )
    p.add_argument(
        "--output-dir",
        type=Path,
        default=Path.home() / "yolo_device_monitor" / "rules",
        help="输出目录，默认 ~/yolo_device_monitor/rules",
    )
    return p.parse_args()


def clean_text(value):
    if value is None:
        return ""
    text = str(value)
    text = text.replace("\xa0", " ")
    text = text.replace("\u3000", " ")
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n+", "\n", text)
    return text.strip()


def compact_text(value):
    return re.sub(r"\s+", "", clean_text(value))


def normalize_header(text):
    return compact_text(text).lower()


def find_column(headers, aliases):
    normalized = [normalize_header(x) for x in headers]
    alias_norm = [normalize_header(x) for x in aliases]

    for i, h in enumerate(normalized):
        if h in alias_norm:
            return i

    for i, h in enumerate(normalized):
        for a in alias_norm:
            if a and a in h:
                return i

    return None


def row_to_texts(row):
    return [clean_text(cell.text) for cell in row.cells]


def detect_header_row(table):
    """
    尝试在表格前几行中寻找表头。
    """
    max_scan = min(len(table.rows), 8)
    best = None
    best_score = -1

    for ridx in range(max_scan):
        cells = row_to_texts(table.rows[ridx])
        score = 0
        for logical_name, aliases in COLUMN_ALIASES.items():
            if find_column(cells, aliases) is not None:
                score += 1

        if score > best_score:
            best_score = score
            best = ridx

    if best_score >= 3:
        return best

    return None


def normalize_state(raw):
    text = compact_text(raw)
    if not text:
        return "unknown"

    if text in STATE_MAP:
        return STATE_MAP[text]

    for k, v in sorted(STATE_MAP.items(), key=lambda x: len(x[0]), reverse=True):
        if k in text:
            return v

    return "unknown"


def normalize_color(raw):
    text = compact_text(raw)
    if not text:
        return "unknown"

    if text in COLOR_MAP:
        return COLOR_MAP[text]

    found = []
    for k, v in COLOR_MAP.items():
        if k in text and v not in found:
            found.append(v)

    if len(found) == 1:
        return found[0]
    if len(found) > 1:
        return "_".join(found)

    return "unknown"


def infer_severity(meaning, normal_state, state):
    text = compact_text(f"{meaning} {normal_state}").lower()

    # 明确“正常状态”字段不直接代表当前这一行一定正常，因此优先看 meaning。
    meaning_text = compact_text(meaning).lower()

    for w in CRITICAL_WORDS:
        if compact_text(w).lower() in meaning_text:
            # “无告警”不能因为包含“告警”被判 critical
            if "无告警" in meaning_text:
                break
            return "critical"

    for w in WARNING_WORDS:
        if compact_text(w).lower() in meaning_text:
            return "warning"

    for w in NORMAL_WORDS:
        if compact_text(w).lower() in meaning_text:
            return "normal"

    for w in INFO_WORDS:
        if compact_text(w).lower() in meaning_text:
            return "info"

    # 辅助规则
    if "无告警" in meaning_text:
        return "normal"

    if state == "off" and any(x in meaning_text for x in ["无供电", "失电", "故障", "down"]):
        return "critical"

    return "unknown"


def split_multivalue(text):
    """
    仅对非常明显的分隔符做切分。
    """
    text = clean_text(text)
    if not text:
        return []

    parts = re.split(r"[；;]+", text)
    return [p.strip() for p in parts if p.strip()]


def extract_rows_from_table(table, table_index):
    header_idx = detect_header_row(table)
    if header_idx is None:
        return [], {
            "table_index": table_index,
            "parsed": False,
            "reason": "未识别到有效表头",
            "row_count": len(table.rows),
        }

    headers = row_to_texts(table.rows[header_idx])

    col_idx = {}
    for logical_name, aliases in COLUMN_ALIASES.items():
        col_idx[logical_name] = find_column(headers, aliases)

    if col_idx["device_name"] is None or col_idx["indicator_name"] is None:
        return [], {
            "table_index": table_index,
            "parsed": False,
            "reason": "缺少设备名称或指示灯列",
            "row_count": len(table.rows),
            "headers": headers,
        }

    last_values = {
        "device_name": "",
        "indicator_name": "",
        "color": "",
        "state": "",
        "meaning": "",
        "normal_state": "",
    }

    extracted = []

    for ridx in range(header_idx + 1, len(table.rows)):
        cells = row_to_texts(table.rows[ridx])

        values = {}
        for logical_name, idx in col_idx.items():
            values[logical_name] = (
                clean_text(cells[idx]) if idx is not None and idx < len(cells) else ""
            )

        # 合并单元格或视觉上空白的情况，向下继承
        for key in ["device_name", "indicator_name", "color"]:
            if values[key]:
                last_values[key] = values[key]
            else:
                values[key] = last_values[key]

        # 正常状态经常只在设备首行出现，也允许继承
        if values["normal_state"]:
            last_values["normal_state"] = values["normal_state"]
        elif last_values["normal_state"]:
            values["normal_state"] = last_values["normal_state"]

        # state / meaning 一般不建议盲目继承
        if not any(values.values()):
            continue

        # 跳过显然是重复表头的行
        if normalize_header(values["device_name"]) in [
            normalize_header(x) for x in COLUMN_ALIASES["device_name"]
        ]:
            continue

        normalized_state = normalize_state(values["state"])
        normalized_color = normalize_color(values["color"])
        severity = infer_severity(
            values["meaning"],
            values["normal_state"],
            normalized_state
        )

        extracted.append({
            "table_index": table_index,
            "source_row_index": ridx + 1,
            "device_name": values["device_name"],
            "indicator_name": values["indicator_name"],
            "color_raw": values["color"],
            "color": normalized_color,
            "state_raw": values["state"],
            "state": normalized_state,
            "meaning": values["meaning"],
            "normal_state": values["normal_state"],
            "severity": severity,
        })

    return extracted, {
        "table_index": table_index,
        "parsed": True,
        "reason": "ok",
        "row_count": len(table.rows),
        "header_row_index": header_idx + 1,
        "headers": headers,
        "extracted_rows": len(extracted),
    }


def build_yaml_structure(df):
    devices = OrderedDict()

    for _, row in df.iterrows():
        device = clean_text(row["device_name"])
        indicator = clean_text(row["indicator_name"])
        state_key = clean_text(row["state"]) or "unknown"

        if not device or not indicator:
            continue

        if device not in devices:
            devices[device] = {
                "normal_state_description": clean_text(row["normal_state"]),
                "indicators": OrderedDict(),
            }

        if not devices[device]["normal_state_description"] and clean_text(row["normal_state"]):
            devices[device]["normal_state_description"] = clean_text(row["normal_state"])

        indicators = devices[device]["indicators"]

        if indicator not in indicators:
            indicators[indicator] = {
                "colors": [],
                "states": OrderedDict(),
            }

        color = clean_text(row["color"])
        color_raw = clean_text(row["color_raw"])

        if color and color != "unknown" and color not in indicators[indicator]["colors"]:
            indicators[indicator]["colors"].append(color)

        state_entry = {
            "color": color,
            "color_raw": color_raw,
            "state_raw": clean_text(row["state_raw"]),
            "meaning": clean_text(row["meaning"]),
            "severity": clean_text(row["severity"]),
        }

        # 同一 indicator + state 可能有多种颜色 / 多行规则
        states = indicators[indicator]["states"]

        if state_key not in states:
            states[state_key] = state_entry
        else:
            existing = states[state_key]
            if isinstance(existing, list):
                existing.append(state_entry)
            else:
                states[state_key] = [existing, state_entry]

    return {
        "schema_version": 1,
        "source_type": "docx_inspection_table",
        "devices": devices,
    }


def write_review_excel(df, path):
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="rules")

        device_summary = (
            df.groupby("device_name")
            .agg(
                indicator_count=("indicator_name", "nunique"),
                rule_count=("indicator_name", "size"),
                critical_count=("severity", lambda x: (x == "critical").sum()),
                warning_count=("severity", lambda x: (x == "warning").sum()),
                unknown_count=("severity", lambda x: (x == "unknown").sum()),
            )
            .reset_index()
        )
        device_summary.to_excel(writer, index=False, sheet_name="device_summary")


def main():
    args = parse_args()

    input_path = args.input.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    output_dir.mkdir(parents=True, exist_ok=True)

    if not input_path.exists():
        raise SystemExit(f"输入文件不存在: {input_path}")

    if input_path.suffix.lower() != ".docx":
        raise SystemExit("当前版本仅支持 .docx 文件")

    print("=" * 90)
    print("巡检表:", input_path)
    print("输出目录:", output_dir)
    print("=" * 90)

    doc = Document(str(input_path))

    all_rows = []
    table_reports = []

    print(f"\n检测到 Word 表格数量: {len(doc.tables)}")

    for i, table in enumerate(doc.tables, 1):
        rows, report = extract_rows_from_table(table, i)
        all_rows.extend(rows)
        table_reports.append(report)

        status = "OK" if report.get("parsed") else "SKIP"
        print(
            f"[{status}] 表格 {i}: "
            f"行数={report.get('row_count', 0)}, "
            f"提取规则={report.get('extracted_rows', 0)}, "
            f"{report.get('reason', '')}"
        )

    if not all_rows:
        raise SystemExit(
            "\n没有解析出任何巡检规则。\n"
            "请把该 Word 文档或表格截图发给我，我再针对实际格式调整解析器。"
        )

    df = pd.DataFrame(all_rows)

    # 基本清洗
    for col in [
        "device_name", "indicator_name", "color_raw", "color",
        "state_raw", "state", "meaning", "normal_state", "severity"
    ]:
        if col in df.columns:
            df[col] = df[col].fillna("").astype(str).map(clean_text)

    # 去掉完全重复规则
    df = df.drop_duplicates(
        subset=[
            "device_name",
            "indicator_name",
            "color_raw",
            "state_raw",
            "meaning",
        ]
    ).reset_index(drop=True)

    csv_path = output_dir / "device_indicator_rules.csv"
    yaml_path = output_dir / "device_indicator_rules.yaml"
    xlsx_path = output_dir / "device_indicator_rules_review.xlsx"
    summary_path = output_dir / "parse_summary.txt"

    df.to_csv(csv_path, index=False, encoding="utf-8-sig")

    yaml_obj = build_yaml_structure(df)
    with open(yaml_path, "w", encoding="utf-8") as f:
        yaml.safe_dump(
            yaml_obj,
            f,
            allow_unicode=True,
            sort_keys=False,
            width=120,
        )

    write_review_excel(df, xlsx_path)

    device_count = int(df["device_name"].nunique())
    indicator_count = int(df[["device_name", "indicator_name"]].drop_duplicates().shape[0])
    critical_count = int((df["severity"] == "critical").sum())
    warning_count = int((df["severity"] == "warning").sum())
    unknown_count = int((df["severity"] == "unknown").sum())

    summary_lines = [
        f"源文件: {input_path}",
        f"Word 表格数量: {len(doc.tables)}",
        f"成功提取规则数: {len(df)}",
        f"设备数量: {device_count}",
        f"设备-指示灯组合数量: {indicator_count}",
        f"critical 规则数: {critical_count}",
        f"warning 规则数: {warning_count}",
        f"unknown 规则数: {unknown_count}",
        "",
        "输出文件:",
        str(csv_path),
        str(yaml_path),
        str(xlsx_path),
        "",
        "注意:",
        "severity 为自动推断结果，首次使用前应人工检查 review.xlsx。",
        "颜色、状态、含义均保留了原始字段，可追溯到 Word 表格。",
    ]

    summary_path.write_text("\n".join(summary_lines), encoding="utf-8")

    print()
    print("=" * 90)
    print("解析完成")
    print("=" * 90)
    print(f"设备数量: {device_count}")
    print(f"规则数量: {len(df)}")
    print(f"设备-指示灯组合: {indicator_count}")
    print(f"critical: {critical_count}")
    print(f"warning: {warning_count}")
    print(f"unknown: {unknown_count}")
    print()
    print("CSV :", csv_path)
    print("YAML:", yaml_path)
    print("复核 :", xlsx_path)
    print("摘要 :", summary_path)
    print()
    print("建议优先打开 device_indicator_rules_review.xlsx 做一次人工检查。")


if __name__ == "__main__":
    main()

